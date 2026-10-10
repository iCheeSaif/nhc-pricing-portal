"""NHC-style independent pricing application. All money is calculated server-side.
No ServiceNow or PPM connection is required to run this standalone app.
"""
from __future__ import annotations

import csv
import io
import json
import os
import secrets
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from typing import Any

from fastapi import FastAPI, Body, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from starlette.middleware.sessions import SessionMiddleware

ROOT = Path(__file__).resolve().parent
ROOT.joinpath('data').mkdir(exist_ok=True)
DB_PATH = Path(os.environ.get('PRICING_DB', ROOT/'data'/'pricing.sqlite3'))
PRODUCTION = os.getenv('APP_ENV', 'demo').lower() == 'production'
SECRET_KEY = os.environ.get('SECRET_KEY', '')
if PRODUCTION and len(SECRET_KEY) < 32:
    raise RuntimeError('SECRET_KEY must contain at least 32 characters in production')
if not SECRET_KEY:
    SECRET_KEY = secrets.token_hex(32)

app = FastAPI(title='بوابة التسعير - شراكات', version='1.0.0', docs_url=None if PRODUCTION else '/api/docs')
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, same_site='lax', https_only=PRODUCTION,
                   max_age=60*60*8, session_cookie='pricing_session')

SCHEMA = '''
CREATE TABLE IF NOT EXISTS users (
 id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE, name TEXT NOT NULL, role TEXT NOT NULL,
 password_salt TEXT NOT NULL, password_hash TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS destinations (
 id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, city TEXT NOT NULL, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS agreements (
 id INTEGER PRIMARY KEY, contract_no TEXT NOT NULL UNIQUE, developer TEXT NOT NULL,
 project_name TEXT NOT NULL, zone TEXT, agreement_model TEXT, start_date TEXT, end_date TEXT,
 land_area REAL DEFAULT 0, meter_price REAL DEFAULT 0, source TEXT NOT NULL DEFAULT 'manual',
 external_ppm_id TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS projects (
 id INTEGER PRIMARY KEY, code TEXT NOT NULL UNIQUE, name TEXT NOT NULL, destination_id INTEGER NOT NULL,
 agreement_id INTEGER NOT NULL UNIQUE, project_type TEXT NOT NULL CHECK(project_type IN ('ministry','company')),
 ministry_factor REAL NOT NULL DEFAULT 0.2941, created_at TEXT NOT NULL,
 FOREIGN KEY(destination_id) REFERENCES destinations(id), FOREIGN KEY(agreement_id) REFERENCES agreements(id)
);
CREATE TABLE IF NOT EXISTS budgets (
 id INTEGER PRIMARY KEY, destination_id INTEGER NOT NULL, name TEXT NOT NULL,
 budget_type TEXT NOT NULL CHECK(budget_type IN ('infra','installment')),
 total_amount REAL NOT NULL CHECK(total_amount >= 0), created_at TEXT NOT NULL,
 FOREIGN KEY(destination_id) REFERENCES destinations(id)
);
CREATE TABLE IF NOT EXISTS units (
 id INTEGER PRIMARY KEY, project_id INTEGER NOT NULL, unit_code TEXT NOT NULL,
 active INTEGER NOT NULL DEFAULT 1, version INTEGER NOT NULL DEFAULT 1,
 data_json TEXT NOT NULL, pricing_json TEXT NOT NULL, last_request_id INTEGER, created_at TEXT NOT NULL,
 UNIQUE(project_id, unit_code), FOREIGN KEY(project_id) REFERENCES projects(id)
);
CREATE TABLE IF NOT EXISTS requests (
 id INTEGER PRIMARY KEY, number TEXT UNIQUE NOT NULL, project_id INTEGER NOT NULL,
 kind TEXT NOT NULL CHECK(kind IN ('base','price','design')), status TEXT NOT NULL DEFAULT 'draft',
 created_by INTEGER NOT NULL, checked_by INTEGER, linked_request_id INTEGER, reasons_json TEXT,
 notes TEXT DEFAULT '', submitted_at TEXT, closed_at TEXT, created_at TEXT NOT NULL,
 FOREIGN KEY(project_id) REFERENCES projects(id), FOREIGN KEY(created_by) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS request_items (
 id INTEGER PRIMARY KEY, request_id INTEGER NOT NULL, unit_id INTEGER,
 unit_code TEXT NOT NULL, operation TEXT NOT NULL DEFAULT 'update',
 before_json TEXT, after_json TEXT, before_pricing TEXT, after_pricing TEXT,
 infra_budget_id INTEGER, installment_budget_id INTEGER,
 UNIQUE(request_id, unit_code), FOREIGN KEY(request_id) REFERENCES requests(id)
);
CREATE TABLE IF NOT EXISTS approvals (
 id INTEGER PRIMARY KEY, request_id INTEGER NOT NULL, step INTEGER NOT NULL, role TEXT NOT NULL,
 user_id INTEGER, status TEXT NOT NULL DEFAULT 'waiting', comment TEXT, acted_at TEXT,
 UNIQUE(request_id, step), FOREIGN KEY(request_id) REFERENCES requests(id)
);
CREATE TABLE IF NOT EXISTS budget_ledger (
 id INTEGER PRIMARY KEY, budget_id INTEGER NOT NULL, request_id INTEGER NOT NULL,
 unit_code TEXT NOT NULL, amount REAL NOT NULL, created_at TEXT NOT NULL,
 FOREIGN KEY(budget_id) REFERENCES budgets(id)
);
CREATE TABLE IF NOT EXISTS audit (
 id INTEGER PRIMARY KEY, entity TEXT NOT NULL, entity_id INTEGER,
 request_id INTEGER, user_id INTEGER, action TEXT NOT NULL, detail TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS unit_history (
 id INTEGER PRIMARY KEY, unit_id INTEGER NOT NULL, request_id INTEGER NOT NULL,
 before_json TEXT, after_json TEXT, before_pricing TEXT, after_pricing TEXT, created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS ix_units_project ON units(project_id,active);
CREATE INDEX IF NOT EXISTS ix_requests_project ON requests(project_id,status);
CREATE INDEX IF NOT EXISTS ix_items_request ON request_items(request_id);
CREATE INDEX IF NOT EXISTS ix_budget_ledger ON budget_ledger(budget_id);
CREATE INDEX IF NOT EXISTS ix_audit_request ON audit(request_id);
'''

ROLES = {'admin':'مدير النظام','maker':'ممثل شراكات - Maker','checker':'مراجع شراكات - Checker',
         'project_manager':'مدير المشروع','pricing_specialist':'مختص التسعير',
         'relationship_manager':'مدير العلاقة','devrel_manager':'مدير علاقات المطورين',
         'operations_manager':'مدير العمليات','pmo_manager':'مدير مكتب إدارة المشاريع','sector_head':'رئيس القطاع'}
APPROVAL_ROLES = ['project_manager','pricing_specialist','relationship_manager','devrel_manager',
                  'operations_manager','pmo_manager','sector_head']
REQUIRED_HEADERS = ['unit_code', 'unit_type','sale_area','land_area','price_method']
HEADERS = ['unit_code','unit_type','status','block','building','floor','model','bedrooms','bathrooms','sale_area',
           'land_area','shared_area','block_land_area','block_total_nsa','commercial_area','price_method',
           'allocation_package','beneficiary_inclusive_rate','beneficiary_building_rate','beneficiary_infra_rate',
           'beneficiary_land_rate','nonbeneficiary_inclusive_rate','nonbeneficiary_building_rate',
           'nonbeneficiary_infra_rate','nonbeneficiary_land_rate','infra_discount','installment_discount',
           'inkind_discount','building_discount','escalation_pct']
ALIASES = {'رقم الوحدة':'unit_code','نوع الوحدة':'unit_type','المساحة البيعية':'sale_area',
           'مساحة الأرض':'land_area','المساحة الصافية':'sale_area', 'رقم البلوك':'block', 'رقم المبنى':'building',
           'رقم الدور':'floor', 'النموذج':'model', 'حالة الوحدة':'status', 'طريقة التسعير':'price_method'}


def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def jdump(x): return json.dumps(x, ensure_ascii=False, separators=(',',':'))
def jload(x): return json.loads(x) if x else None

@contextmanager
def db(write=False):
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory=sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    conn.execute('PRAGMA busy_timeout=30000')
    if write: conn.execute('BEGIN IMMEDIATE')
    try:
        yield conn
        if write: conn.commit()
    except:
        if write: conn.rollback()
        raise
    finally: conn.close()

def sha_password(password,salt):
    import hashlib
    return hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt),350_000).hex()
def make_user(conn,username,name,role,password):
    salt=secrets.token_hex(16)
    conn.execute('INSERT INTO users(username,name,role,password_salt,password_hash) VALUES(?,?,?,?,?)',
                 (username,name,role,salt,sha_password(password,salt)))

def audit(conn, entity, entity_id, request_id, user_id, action, detail=''):
    conn.execute('INSERT INTO audit(entity,entity_id,request_id,user_id,action,detail,created_at) VALUES(?,?,?,?,?,?,?)',
                 (entity,entity_id,request_id,user_id,action,detail,now()))

def init_db():
    with db(write=True) as c:
        c.executescript(SCHEMA)
        if c.execute('SELECT COUNT(*) n FROM users').fetchone()['n']:
            return
        password = os.environ.get('INITIAL_ADMIN_PASSWORD', 'Demo@2026!')
        if PRODUCTION and not os.environ.get('INITIAL_ADMIN_PASSWORD'):
            raise RuntimeError('INITIAL_ADMIN_PASSWORD is mandatory in production on first boot')
        make_user(c,'admin','مدير النظام','admin',password)
        if PRODUCTION:
            return
        demos=[('maker','أحمد - شراكات','maker'),('checker','سارة - المراجع','checker'),
               ('pm','مدير المشروع','project_manager'),('pricing','مختص التسعير','pricing_specialist'),
               ('relation','مدير العلاقة','relationship_manager'),('developers','مدير علاقات المطورين','devrel_manager'),
               ('ops','مدير العمليات','operations_manager'),('pmo','مدير PMO','pmo_manager'),
               ('head','رئيس القطاع','sector_head')]
        for user,name,role in demos:make_user(c,user,name,role,password)
        timestamp=now()
        c.execute('INSERT INTO destinations(name,city,created_at) VALUES(?,?,?)',('وجهة النخيل','الرياض',timestamp))
        dest=c.execute('SELECT last_insert_rowid() id').fetchone()['id']
        c.execute('''INSERT INTO budgets(destination_id,name,budget_type,total_amount,created_at) VALUES(?,?,?,?,?)''',
                  (dest,'اتفاقية بنية تحتية - 2026','infra',6_000_000,timestamp))
        c.execute('''INSERT INTO budgets(destination_id,name,budget_type,total_amount,created_at) VALUES(?,?,?,?,?)''',
                  (dest,'اتفاقية القسط الميسر - 2026','installment',4_000_000,timestamp))
        c.execute('''INSERT INTO agreements(contract_no,developer,project_name,zone,agreement_model,start_date,end_date,land_area,meter_price,source,created_at)
                     VALUES(?,?,?,?,?,?,?,?,?,?,?)''',('DEV-2026-001','شركة تطوير النخيل','حي النخيل السكني','Z-01','اتفاقية تطوير','2026-01-01','2028-12-31',120000,4200,'sample',timestamp))
        agreement=c.execute('SELECT last_insert_rowid() id').fetchone()['id']
        c.execute('INSERT INTO projects(code,name,destination_id,agreement_id,project_type,created_at) VALUES(?,?,?,?,?,?)',
                  ('PRJ-001','حي النخيل السكني',dest,agreement,'ministry',timestamp))
        pid=c.execute('SELECT last_insert_rowid() id').fetchone()['id']
        for i in range(1,13):
            snap=default_snapshot({'unit_code':f'V-{i:03d}','unit_type':'villa','status':'available',
                                  'block':'B-01' if i<=6 else 'B-02','sale_area':210+i*5,'land_area':190+i*4,
                                  'price_method':'detailed','beneficiary_building_rate':1850,
                                  'beneficiary_infra_rate':480,'beneficiary_land_rate':620,
                                  'nonbeneficiary_building_rate':2000,'nonbeneficiary_infra_rate':530,
                                  'nonbeneficiary_land_rate':740})
            calc=calculate(snap,{'project_type':'ministry','ministry_factor':0.2941})
            c.execute('INSERT INTO units(project_id,unit_code,data_json,pricing_json,created_at) VALUES(?,?,?,?,?)',
                      (pid,snap['unit_code'],jdump(snap),jdump(calc),timestamp))


def require_user(request: Request):
    uid=request.session.get('uid')
    if not uid: raise HTTPException(401,'الرجاء تسجيل الدخول')
    with db() as c: row=c.execute('SELECT id,username,name,role,active FROM users WHERE id=?',(uid,)).fetchone()
    if not row or not row['active']:raise HTTPException(401,'الجلسة غير صالحة')
    if request.method not in ('GET','HEAD','OPTIONS'):
        if request.headers.get('x-csrf-token') != request.session.get('csrf'):
            raise HTTPException(403,'رمز الحماية غير صالح')
    return dict(row)

def authorize(user,roles):
    if user['role'] not in roles and user['role']!='admin':raise HTTPException(403,'ليست لديك صلاحية لتنفيذ هذه العملية')
def must(c,table,id):
    if table not in ('projects','agreements','destinations','budgets','requests','units'):raise ValueError('bad table')
    row=c.execute(f'SELECT * FROM {table} WHERE id=?',(id,)).fetchone()
    if not row:raise HTTPException(404,'السجل غير موجود')
    return row

def D(v,nonnegative=True):
    try:
        out=Decimal(str(0 if v is None or v=='' else v))
    except (InvalidOperation,ValueError): raise ValueError('يجب إدخال قيمة رقمية صحيحة')
    if not out.is_finite() or (nonnegative and out<0):raise ValueError('القيمة يجب أن تكون رقمًا غير سالب')
    return out

def money(v): return float(D(v,False).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))
def qty(v): return float(D(v).quantize(Decimal('0.0001'),rounding=ROUND_HALF_UP))


def default_snapshot(raw,old=None):
    """Whitelisted, stable unit JSON. Unknown Excel columns are never trusted."""
    snapshot=dict(old or {})
    text_fields=['unit_code','unit_type','status','block','building','floor','model','price_method','allocation_package']
    numeric_fields=['bedrooms','bathrooms','sale_area','land_area','shared_area','block_land_area',
                    'block_total_nsa','commercial_area','beneficiary_inclusive_rate','beneficiary_building_rate',
                    'beneficiary_infra_rate','beneficiary_land_rate','nonbeneficiary_inclusive_rate',
                    'nonbeneficiary_building_rate','nonbeneficiary_infra_rate','nonbeneficiary_land_rate',
                    'infra_discount','installment_discount','inkind_discount','building_discount','escalation_pct']
    for f in text_fields:
        if f in raw and raw[f] is not None:snapshot[f]=str(raw[f]).strip()
    for f in numeric_fields:
        if f in raw and raw[f] is not None and str(raw[f]).strip()!='':snapshot[f]=qty(raw[f])
    for f in text_fields:snapshot.setdefault(f,'')
    for f in numeric_fields:snapshot.setdefault(f,0)
    snapshot['unit_code']=str(snapshot.get('unit_code','')).strip()
    if not snapshot['unit_code']:raise ValueError('رقم الوحدة مطلوب')
    if len(snapshot['unit_code'])>60:raise ValueError('رقم الوحدة طويل جدًا')
    if snapshot['unit_type'] not in ('villa','townhouse','apartment'):raise ValueError('نوع الوحدة غير صالح')
    if snapshot['status'] not in ('preliminary','available','reserved','sold'):
        raise ValueError('حالة الوحدة غير صحيحة')
    if snapshot['price_method'] not in ('inclusive','detailed','allocated'):
        raise ValueError('طريقة التسعير غير صالحة')
    if snapshot['price_method']=='allocated' and snapshot['allocation_package'] not in ('building_infra','building_land','infra_land'):
        raise ValueError('الرجاء اختيار حزمة التخصيص')
    if snapshot['sale_area']<=0:raise ValueError('المساحة البيعية يجب أن تكون أكبر من صفر')
    if snapshot['unit_type']=='apartment' and snapshot['block_total_nsa']<=0:
        raise ValueError('الشقق تتطلب إدخال مجموع المساحات الصافية للبلوك')
    return snapshot


def land_basis(s,p):
    if s['unit_type']!='apartment':return D(s['land_area'])
    block_area=D(s['block_land_area'])
    block_nsa=D(s['block_total_nsa'])
    common=(D(s['sale_area'])/block_nsa)*block_area
    if p['project_type']=='ministry':
        # BRD field 111: (unit shared allocated area + common land share) * 0.2941.
        return (common+D(s['shared_area']))*D(p['ministry_factor'])
    commercial=D(s['commercial_area'])
    return common * (block_nsa/(block_nsa+commercial)) if commercial else common


def calculate(s,p):
    """Central deterministic pricing engine; each segment uses the same formula.
    BRD-supported basics plus explicit provisional assumptions documented in README.
    """
    area=D(s['sale_area']); land_area=land_basis(s,p)
    discount_infra=D(s['infra_discount']); discount_installment=D(s['installment_discount'])
    discount_inkind=D(s['inkind_discount']); discount_building=D(s['building_discount'])
    escalation=D(s['escalation_pct'])
    if escalation>100:raise ValueError('نسبة التدرج تتجاوز 100%')
    if (discount_infra or discount_installment) and escalation:
        raise ValueError('لا يسمح بالتدرج مع خصم البنية التحتية أو القسط الميسر')
    if discount_inkind and p['project_type']!='ministry':
        raise ValueError('الدعم العيني مخصص لمشاريع أراضي الوزارة')
    if (discount_infra or discount_installment) and s['unit_type']=='apartment' and land_area<=0:
        raise ValueError('تعذر احتساب حصة الأرض للشقة')
    out={'land_basis':money(land_area), 'price_method':s['price_method'],'segments':{}}
    for segment in ('beneficiary','nonbeneficiary'):
        inc=D(s[f'{segment}_inclusive_rate'])
        br=D(s[f'{segment}_building_rate']); ir=D(s[f'{segment}_infra_rate']); lr=D(s[f'{segment}_land_rate'])
        land=land_area*lr; infra=land_area*ir; building=area*br
        if s['price_method']=='inclusive':
            total=area*inc
            building=total-infra-land
            if building<0:raise ValueError(f'السعر الشامل لفئة {segment} أقل من مجموع الأرض والبنية التحتية')
        elif s['price_method']=='allocated':
            pkg=s['allocation_package']
            if pkg=='building_infra':building=area*br; infra=0
            elif pkg=='building_land':building=area*br;land=0
            elif pkg=='infra_land':infra=land_area*ir;land=0
        if segment=='beneficiary':
            if discount_infra>infra:raise ValueError('خصم البنية التحتية يتجاوز قيمة البنية التحتية')
            if discount_installment + discount_building>building:raise ValueError('خصم البناء والقسط الميسر يتجاوز قيمة البناء')
            if discount_inkind>land:raise ValueError('خصم الدعم العيني يتجاوز قيمة الأرض الخام')
            infra-=discount_infra; building-=discount_installment+discount_building;land-=discount_inkind
        subtotal=building+infra+land
        increase=subtotal*escalation/100
        total=subtotal+increase
        if segment=='beneficiary' and discount_installment>0 and total<250000:
            raise ValueError('السعر النهائي بعد خصم القسط الميسر يجب ألا يقل عن 250,000 ريال')
        out['segments'][segment]={
            'building':money(building),'infra':money(infra),'land':money(land),
            'escalation':money(increase),'total':money(total)}
    return out

def load_unit(row):
    r=dict(row)
    r['data']=jload(r.pop('data_json'))
    r['pricing']=jload(r.pop('pricing_json'))
    return r

def full_request(c,row):
    r=dict(row)
    r['reasons']=jload(r.get('reasons_json')) or []
    r['items']=[{**dict(i),'before':jload(i['before_json']),'after':jload(i['after_json']),
                 'before_pricing':jload(i['before_pricing']),'after_pricing':jload(i['after_pricing'])}
                for i in c.execute('SELECT * FROM request_items WHERE request_id=? ORDER BY unit_code',(r['id'],))]
    r['approvals']=[dict(x) for x in c.execute('SELECT * FROM approvals WHERE request_id=? ORDER BY step',(r['id'],))]
    r['audit']=[dict(x) for x in c.execute('SELECT a.*,u.name actor FROM audit a LEFT JOIN users u ON a.user_id=u.id WHERE request_id=? ORDER BY a.id DESC',(r['id'],))]
    return r


def make_item(c,req,p,raw,operation='update',infra_budget_id=None,installment_budget_id=None):
    code=str(raw.get('unit_code','')).strip()
    existing=c.execute('SELECT * FROM units WHERE project_id=? AND unit_code=? AND active=1',(p['id'],code)).fetchone()
    old=jload(existing['data_json']) if existing else None
    old_price=jload(existing['pricing_json']) if existing else None
    if req['kind']=='base' and existing:
        raise ValueError(f'الوحدة {code} مسعّرة مسبقًا؛ استخدم طلب تغيير أسعار')
    if req['kind']!='base' and not existing and operation!='add':
        raise ValueError(f'الوحدة {code} غير موجودة')
    if req['kind']=='price' and not existing:raise ValueError('لا يمكن إضافة وحدة بواسطة طلب تغيير أسعار')
    if existing and old['status'] in ('reserved','sold'):
        raise ValueError(f'الوحدة {code} محجوزة أو مباعة ولا يمكن تعديلها')
    if operation=='remove' and req['kind']!='design':raise ValueError('الحذف مسموح فقط في تغيير التصاميم')
    if operation=='remove' and not existing:raise ValueError('لا يمكن حذف وحدة غير موجودة')
    if req['kind']=='design' and existing:
        fixed_price_fields=[f'{seg}_{k}' for seg in ('beneficiary','nonbeneficiary') for k in ('inclusive_rate','building_rate','infra_rate','land_rate')]
        fixed_price_fields+=['infra_discount','installment_discount','inkind_discount','building_discount','escalation_pct']
        if any(f in raw and qty(raw[f]) != old.get(f,0) for f in fixed_price_fields):
            raise ValueError('تعديل معدلات الأسعار أو الخصومات يتم عبر طلب تغيير أسعار مستقل')
        if 'status' in raw and raw['status']!=old['status']:
            raise ValueError('حالة الوحدة لا تتغير عبر تعديل التصميم')
        # Existing price inputs/discounts are inherited to recompute on the changed design.
        for key in ('price_method','allocation_package'):
            if key in raw and raw[key]!=old.get(key):raise ValueError('تغيير طريقة التسعير يتم عبر طلب الأسعار')
    after=None if operation=='remove' else default_snapshot(raw,old)
    if after and after['unit_code']!=code:raise ValueError('لا يمكن تغيير معرف الوحدة ضمن الطلب')
    if req['kind']=='price' and after:
        protected=['status','unit_type','sale_area','land_area','shared_area','block','building','floor','model',
                   'block_land_area','block_total_nsa','commercial_area']
        if any(after.get(f)!=old.get(f) for f in protected):
            raise ValueError('تغيير بيانات التصميم يجب أن يتم عبر طلب تغيير التصاميم')
    after_price=calculate(after,p) if after else None
    c.execute('''INSERT INTO request_items(request_id,unit_id,unit_code,operation,before_json,after_json,before_pricing,after_pricing,infra_budget_id,installment_budget_id)
                 VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(request_id,unit_code) DO UPDATE SET
                 operation=excluded.operation,before_json=excluded.before_json,after_json=excluded.after_json,
                 before_pricing=excluded.before_pricing,after_pricing=excluded.after_pricing,
                 infra_budget_id=excluded.infra_budget_id,installment_budget_id=excluded.installment_budget_id''',
              (req['id'],existing['id'] if existing else None,code,operation,jdump(old) if old else None,
               jdump(after) if after else None,jdump(old_price) if old_price else None,
               jdump(after_price) if after_price else None,infra_budget_id,installment_budget_id))
    return after_price

def validate_budget_choice(c,budget_id,project_id,expected):
    if budget_id is None or str(budget_id)=='':return None
    b=must(c,'budgets',int(budget_id))
    p=must(c,'projects',project_id)
    if b['destination_id']!=p['destination_id'] or b['budget_type']!=expected:
        raise ValueError('اتفاقية الميزانية لا تنتمي إلى الوجهة أو نوع الخصم')
    return int(budget_id)

def budget_changes(c, req):
    amounts={}
    items=c.execute('SELECT * FROM request_items WHERE request_id=?',(req['id'],)).fetchall()
    for item in items:
        new=jload(item['after_json']) or {};old=jload(item['before_json']) or {}
        for typ,key,field in [('infra','infra_discount','infra_budget_id'),('installment','installment_discount','installment_budget_id')]:
            delta=D(new.get(key,0),False)-D(old.get(key,0),False)
            if not delta:continue
            bid=item[field]
            if delta>0 and not bid:
                raise ValueError(f'يجب اختيار اتفاقية ميزانية خصم {typ} للوحدة {item["unit_code"]}')
            if delta<0 and not bid:
                raise ValueError(f'يجب اختيار اتفاقية الميزانية المرتبطة بالخصم السابق للوحدة {item["unit_code"]}')
            amounts.setdefault(bid,[]).append((item['unit_code'],money(delta)))
    return amounts

def check_request(c,req,p):
    items=c.execute('SELECT * FROM request_items WHERE request_id=?',(req['id'],)).fetchall()
    if not items:raise ValueError('لا يمكن إرسال طلب دون وحدات')
    warnings=[]
    for item in items:
        existing=c.execute('SELECT * FROM units WHERE project_id=? AND unit_code=? AND active=1',(p['id'],item['unit_code'])).fetchone()
        if item['before_json']:
            if not existing or jload(existing['data_json'])!=jload(item['before_json']):
                raise ValueError(f'تغيرت بيانات الوحدة {item["unit_code"]} بعد تجهيز الطلب، يرجى تحديثه')
            if jload(existing['data_json'])['status'] in ('reserved','sold'):
                raise ValueError(f'الوحدة {item["unit_code"]} مباعة أو محجوزة')
        elif existing:
            raise ValueError(f'تم إنشاء الوحدة {item["unit_code"]} بواسطة طلب آخر')
    if p['project_type']=='ministry':
        current={x['unit_code']:jload(x['data_json']) for x in c.execute('SELECT * FROM units WHERE project_id=? AND active=1',(p['id'],))}
        for item in items:
            nxt=jload(item['after_json'])
            if nxt:current[item['unit_code']]=nxt
            else:current.pop(item['unit_code'],None)
        # Applies only to units explicitly offered at non-beneficiary rate.
        nonbenef=[u for u in current.values() if u['nonbeneficiary_inclusive_rate']>0 or
                 u['nonbeneficiary_building_rate']>0 or u['nonbeneficiary_infra_rate']>0 or
                 u['nonbeneficiary_land_rate']>0]
        if current and len(nonbenef)>len(current)*0.25:
            warnings.append('تنبيه: عدد الوحدات المسعرة لغير المستفيدين يتجاوز 25%؛ يلزم توثيق الاستثناء/تأكيد تعريف النسبة')
    for bid,values in budget_changes(c,req).items():
        b=must(c,'budgets',bid)
        spent=D(c.execute('SELECT COALESCE(SUM(amount),0) n FROM budget_ledger WHERE budget_id=?',(bid,)).fetchone()['n'],False)
        delta=sum((D(v,False) for code,v in values),Decimal(0))
        if spent+delta>D(b['total_amount']):raise ValueError(f'طلب الخصومات يتجاوز ميزانية {b["name"]}')
        if spent+delta >= D(b['total_amount'])*Decimal('.9'):
            warnings.append(f'تنبيه: استهلاك 90% أو أكثر من ميزانية {b["name"]}')
    return warnings

def apply_request(c,req,actor):
    p=must(c,'projects',req['project_id'])
    warnings=check_request(c,req,p)
    items=c.execute('SELECT * FROM request_items WHERE request_id=? ORDER BY id',(req['id'],)).fetchall()
    for item in items:
        old=jload(item['before_json']); new=jload(item['after_json']); price=jload(item['after_pricing'])
        if item['operation']=='remove':
            c.execute('UPDATE units SET active=0,version=version+1,last_request_id=? WHERE id=?',(req['id'],item['unit_id']))
            unit_id=item['unit_id']
        elif item['unit_id']:
            c.execute('UPDATE units SET active=1,version=version+1,data_json=?,pricing_json=?,last_request_id=? WHERE id=?',
                      (jdump(new),jdump(price),req['id'],item['unit_id']))
            unit_id=item['unit_id']
        else:
            old_deleted=c.execute('SELECT id FROM units WHERE project_id=? AND unit_code=? AND active=0',(p['id'],item['unit_code'])).fetchone()
            if old_deleted:
                unit_id=old_deleted['id']
                c.execute('UPDATE units SET active=1,data_json=?,pricing_json=?,version=version+1,last_request_id=? WHERE id=?',
                          (jdump(new),jdump(price),req['id'],unit_id))
            else:
                c.execute('INSERT INTO units(project_id,unit_code,data_json,pricing_json,last_request_id,created_at) VALUES(?,?,?,?,?,?)',
                          (p['id'],item['unit_code'],jdump(new),jdump(price),req['id'],now()))
                unit_id=c.execute('SELECT last_insert_rowid() id').fetchone()['id']
        c.execute('INSERT INTO unit_history(unit_id,request_id,before_json,after_json,before_pricing,after_pricing,created_at) VALUES(?,?,?,?,?,?,?)',
                  (unit_id,req['id'],jdump(old) if old else None,jdump(new) if new else None,
                   item['before_pricing'],item['after_pricing'],now()))
    for bid,changes in budget_changes(c,req).items():
        for code,delta in changes:
            c.execute('INSERT INTO budget_ledger(budget_id,request_id,unit_code,amount,created_at) VALUES(?,?,?,?,?)',
                      (bid,req['id'],code,delta,now()))
    c.execute('UPDATE requests SET status="approved",closed_at=? WHERE id=?',(now(),req['id']))
    audit(c,'request',req['id'],req['id'],actor['id'],'approved',jdump({'warnings':warnings}))
    return warnings

@app.exception_handler(ValueError)
async def validation_exception(request,exc):return JSONResponse(status_code=422,content={'detail':str(exc)})

@app.get('/')
def home():return FileResponse(ROOT/'static'/'index.html')
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')

@app.post('/api/auth/login')
def login(request:Request,payload:dict=Body(...)):
    name=str(payload.get('username','')).strip().lower()
    password=str(payload.get('password',''))
    with db() as c:row=c.execute('SELECT * FROM users WHERE username=? AND active=1',(name,)).fetchone()
    import hmac
    if not row or not hmac.compare_digest(sha_password(password,row['password_salt']),row['password_hash']):
        raise HTTPException(401,'اسم المستخدم أو كلمة المرور غير صحيحة')
    request.session.clear()
    request.session['uid']=row['id'];request.session['csrf']=secrets.token_urlsafe(24)
    return {'user':{k:row[k] for k in ('id','name','username','role')},'csrf':request.session['csrf']}

@app.get('/api/config')
def config():return {'demo':not PRODUCTION,'version':'1.0.0'}

@app.get('/api/auth/me')
def me(request:Request,user=Depends(require_user)):
    return {'user':user,'csrf':request.session['csrf'],'roles':ROLES,'demo':not PRODUCTION}

@app.post('/api/auth/logout')
def logout(request:Request,user=Depends(require_user)):
    request.session.clear();return {'ok':True}

@app.get('/api/users')
def users(user=Depends(require_user)):
    authorize(user,['admin'])
    with db() as c:return [dict(x) for x in c.execute('SELECT id,username,name,role,active FROM users ORDER BY id')]

@app.post('/api/users')
def add_user(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['admin'])
    role=p.get('role')
    if role not in ROLES:raise ValueError('الدور غير صالح')
    username=str(p.get('username','')).strip().lower()
    password=str(p.get('password',''))
    if not username or not str(p.get('name','')).strip():raise ValueError('اسم المستخدم والاسم الكامل مطلوبان')
    if len(password)<10:raise ValueError('كلمة المرور لا تقل عن 10 أحرف')
    with db(write=True) as c:
        try:make_user(c,username,str(p.get('name','')).strip(),role,password)
        except sqlite3.IntegrityError:raise ValueError('اسم المستخدم موجود مسبقًا')
        uid=c.execute('SELECT last_insert_rowid() id').fetchone()['id']
        audit(c,'user',uid,None,user['id'],'created',username)
    return {'id':uid}

@app.get('/api/destinations')
def destinations(user=Depends(require_user)):
    with db() as c:
        result=[]
        for x in c.execute('SELECT * FROM destinations ORDER BY id DESC'):
            r=dict(x);r['project_count']=c.execute('SELECT COUNT(*) n FROM projects WHERE destination_id=?',(x['id'],)).fetchone()['n']
            result.append(r)
        return result

@app.post('/api/destinations')
def add_destination(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    if not str(p.get('name','')).strip() or not str(p.get('city','')).strip():raise ValueError('اسم الوجهة والمدينة مطلوبان')
    with db(write=True) as c:
        try:
            cur=c.execute('INSERT INTO destinations(name,city,created_at) VALUES(?,?,?)',
                          (p['name'].strip(),p['city'].strip(),now()))
        except sqlite3.IntegrityError:raise ValueError('الوجهة موجودة مسبقًا')
        audit(c,'destination',cur.lastrowid,None,user['id'],'created')
        return {'id':cur.lastrowid}

@app.patch('/api/destinations/{destination_id}')
def edit_destination(destination_id:int,p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    name=str(p.get('name','')).strip();city=str(p.get('city','')).strip()
    if not name or not city:raise ValueError('اسم الوجهة والمدينة مطلوبان')
    with db(write=True) as c:
        must(c,'destinations',destination_id)
        try:c.execute('UPDATE destinations SET name=?,city=? WHERE id=?',(name,city,destination_id))
        except sqlite3.IntegrityError:raise ValueError('هذا الاسم مستخدم لوجهة أخرى')
        audit(c,'destination',destination_id,None,user['id'],'updated')
    return {'ok':True}

@app.get('/api/budgets')
def budgets(destination_id:int|None=None,user=Depends(require_user)):
    with db() as c:
        sql='''SELECT b.*,COALESCE((SELECT SUM(l.amount) FROM budget_ledger l WHERE l.budget_id=b.id),0) used
                  FROM budgets b WHERE (? IS NULL OR destination_id=?) ORDER BY id DESC'''
        result=[]
        for row in c.execute(sql,(destination_id,destination_id)):
            r=dict(row);r['remaining']=money(D(r['total_amount'])-D(r['used'],False))
            r['usage_pct']=round(100*r['used']/r['total_amount'],1) if r['total_amount'] else 0
            result.append(r)
        return result

@app.post('/api/budgets')
def add_budget(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    t=p.get('budget_type')
    if t not in ('infra','installment'):raise ValueError('نوع الميزانية غير صالح')
    with db(write=True) as c:
        must(c,'destinations',p.get('destination_id'))
        cur=c.execute('INSERT INTO budgets(destination_id,name,budget_type,total_amount,created_at) VALUES(?,?,?,?,?)',
                     (p['destination_id'],str(p.get('name','')).strip(),t,money(p.get('total_amount')),now()))
        audit(c,'budget',cur.lastrowid,None,user['id'],'created')
        return {'id':cur.lastrowid}

@app.patch('/api/budgets/{budget_id}')
def edit_budget(budget_id:int,p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    with db(write=True) as c:
        budget=must(c,'budgets',budget_id)
        total=money(p.get('total_amount',budget['total_amount']))
        spent=D(c.execute('SELECT COALESCE(SUM(amount),0) n FROM budget_ledger WHERE budget_id=?',(budget_id,)).fetchone()['n'],False)
        if D(total)<spent:raise ValueError('الميزانية الجديدة أقل من إجمالي الخصومات المعتمدة')
        name=str(p.get('name',budget['name'])).strip()
        if not name:raise ValueError('اسم الميزانية مطلوب')
        c.execute('UPDATE budgets SET name=?,total_amount=? WHERE id=?',(name,total,budget_id))
        audit(c,'budget',budget_id,None,user['id'],'updated',f'{budget["total_amount"]} → {total}')
    return {'ok':True}

@app.get('/api/agreements')
def agreements(user=Depends(require_user)):
    with db() as c:return [dict(x) for x in c.execute('SELECT * FROM agreements ORDER BY id DESC')]

@app.post('/api/agreements')
def add_agreement(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    fields=['contract_no','developer','project_name','zone','agreement_model','start_date','end_date','land_area','meter_price']
    if not all(str(p.get(f,'')).strip() for f in fields[:3]):raise ValueError('رقم العقد والمطور والمشروع مطلوبة')
    vals=[str(p.get(f,'')).strip() for f in fields[:7]]+[money(p.get('land_area',0)),money(p.get('meter_price',0))]
    with db(write=True) as c:
        try:cur=c.execute('''INSERT INTO agreements(contract_no,developer,project_name,zone,agreement_model,start_date,end_date,land_area,meter_price,created_at)
                             VALUES(?,?,?,?,?,?,?,?,?,?)''',(*vals,now()))
        except sqlite3.IntegrityError:raise ValueError('رقم العقد موجود مسبقًا')
        audit(c,'agreement',cur.lastrowid,None,user['id'],'created')
        return {'id':cur.lastrowid}

@app.get('/api/projects')
def projects(user=Depends(require_user)):
    with db() as c:
        return [dict(x) for x in c.execute('''SELECT p.*,d.name destination_name,a.contract_no,a.developer,
                   (SELECT COUNT(*) FROM units u WHERE u.project_id=p.id AND u.active=1) unit_count
                   FROM projects p JOIN destinations d ON d.id=p.destination_id JOIN agreements a ON a.id=p.agreement_id ORDER BY p.id DESC''')]

@app.post('/api/projects')
def add_project(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    if p.get('project_type') not in ('ministry','company'):raise ValueError('نوع المشروع غير صالح')
    with db(write=True) as c:
        must(c,'destinations',p.get('destination_id'));must(c,'agreements',p.get('agreement_id'))
        try:
            cur=c.execute('INSERT INTO projects(code,name,destination_id,agreement_id,project_type,ministry_factor,created_at) VALUES(?,?,?,?,?,?,?)',
                          (str(p['code']).strip(),str(p['name']).strip(),p['destination_id'],p['agreement_id'],p['project_type'],
                           qty(p.get('ministry_factor',0.2941)),now()))
        except sqlite3.IntegrityError:raise ValueError('كود المشروع مستخدم أو الاتفاقية مرتبطة بمشروع آخر')
        audit(c,'project',cur.lastrowid,None,user['id'],'created');return {'id':cur.lastrowid}

@app.get('/api/units')
def units(project_id:int,user=Depends(require_user)):
    with db() as c:
        must(c,'projects',project_id)
        return [load_unit(r) for r in c.execute('SELECT * FROM units WHERE project_id=? AND active=1 ORDER BY unit_code',(project_id,))]

@app.get('/api/units/{unit_id}/history')
def history(unit_id:int,user=Depends(require_user)):
    with db() as c:
        must(c,'units',unit_id)
        return [dict(x) for x in c.execute('SELECT * FROM unit_history WHERE unit_id=? ORDER BY id DESC',(unit_id,))]

@app.post('/api/units/{unit_id}/status')
def unit_status(unit_id:int,p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    status=p.get('status')
    if status not in ('preliminary','available','reserved','sold'):raise ValueError('حالة غير صالحة')
    with db(write=True) as c:
        row=must(c,'units',unit_id)
        if not row['active']:raise ValueError('الوحدة محذوفة')
        old=jload(row['data_json']);new=dict(old)
        if old['status']=='sold' and status!='sold':raise ValueError('الوحدة المباعة لا يمكن إرجاعها للحالة السابقة')
        new['status']=status
        c.execute('UPDATE units SET data_json=?,version=version+1 WHERE id=?',(jdump(new),unit_id))
        audit(c,'unit',unit_id,None,user['id'],'status_changed',f'{old["status"]} → {status}')
        return {'ok':True}

@app.get('/api/requests')
def requests(kind:str|None=None,status:str|None=None,project_id:int|None=None,user=Depends(require_user)):
    with db() as c:
        return [dict(x) for x in c.execute('''SELECT r.*,p.name project_name,u.name creator_name,
                          (SELECT COUNT(*) FROM request_items WHERE request_id=r.id) item_count
                          FROM requests r JOIN projects p ON p.id=r.project_id JOIN users u ON u.id=r.created_by
                          WHERE (? IS NULL OR r.kind=?) AND (? IS NULL OR r.status=?) AND (? IS NULL OR r.project_id=?)
                          ORDER BY r.id DESC''',(kind,kind,status,status,project_id,project_id))]

@app.post('/api/requests')
def add_request(p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    if p.get('kind') not in ('base','price','design'):raise ValueError('نوع الطلب غير صالح')
    with db(write=True) as c:
        project=must(c,'projects',p.get('project_id'))
        if p['kind']!='base' and not c.execute('SELECT 1 FROM units WHERE project_id=? AND active=1',(project['id'],)).fetchone():
            raise ValueError('يجب تعريف التسعير المبدئي أولًا')
        linked=p.get('linked_request_id')
        if linked:
            linked_req=must(c,'requests',linked)
            if linked_req['project_id']!=project['id'] or linked_req['status']!='approved':
                raise ValueError('الطلب المرتبط غير معتمد أو يتبع مشروعًا آخر')
        reasons=p.get('reasons',[])
        if not isinstance(reasons,list) or len(reasons)>12 or any(not isinstance(k,str) or len(k)>60 for k in reasons):
            raise ValueError('قائمة أسباب تغيير الأسعار غير صحيحة')
        cur=c.execute('INSERT INTO requests(number,project_id,kind,created_by,notes,linked_request_id,reasons_json,created_at) VALUES(?,?,?,?,?,?,?,?)',
                      (f'TEMP-{secrets.token_hex(5)}',project['id'],p['kind'],user['id'],str(p.get('notes','')),linked,jdump(reasons),now()))
        code={'base':'BAS','price':'PRC','design':'DSN'}[p['kind']]
        number=f'{code}{cur.lastrowid:06d}'
        c.execute('UPDATE requests SET number=? WHERE id=?',(number,cur.lastrowid))
        audit(c,'request',cur.lastrowid,cur.lastrowid,user['id'],'created')
        return {'id':cur.lastrowid,'number':number}

@app.get('/api/requests/{request_id}')
def request_details(request_id:int,user=Depends(require_user)):
    with db() as c:return full_request(c,must(c,'requests',request_id))

@app.post('/api/requests/{request_id}/items')
def save_item(request_id:int,p:dict=Body(...),user=Depends(require_user)):
    authorize(user,['maker'])
    with db(write=True) as c:
        req=must(c,'requests',request_id)
        if req['status'] not in ('draft','rework'):raise ValueError('الطلب لا يقبل التعديل في حالته الحالية')
        if req['created_by']!=user['id'] and user['role']!='admin':raise HTTPException(403,'يسمح بالتعديل لمنشئ الطلب فقط')
        proj=must(c,'projects',req['project_id'])
        infra=validate_budget_choice(c,p.get('infra_budget_id'),proj['id'],'infra')
        installment=validate_budget_choice(c,p.get('installment_budget_id'),proj['id'],'installment')
        raw=p.get('data')
        if not isinstance(raw,dict):raise ValueError('بيانات الوحدة غير صالحة')
        price=make_item(c,req,proj,raw,p.get('operation','update'),infra,installment)
        audit(c,'request',request_id,request_id,user['id'],'item_saved',str(raw.get('unit_code')))
        return {'ok':True,'pricing':price}

@app.delete('/api/requests/{request_id}/items/{item_id}')
def delete_item(request_id:int,item_id:int,user=Depends(require_user)):
    authorize(user,['maker'])
    with db(write=True) as c:
        req=must(c,'requests',request_id)
        if req['status'] not in ('draft','rework'):raise ValueError('الطلب غير قابل للتعديل')
        if req['created_by']!=user['id'] and user['role']!='admin':raise HTTPException(403,'لا تملك صلاحية الحذف')
        c.execute('DELETE FROM request_items WHERE id=? AND request_id=?',(item_id,request_id))
        audit(c,'request',request_id,request_id,user['id'],'item_deleted',str(item_id))
        return {'ok':True}

@app.get('/api/templates/units.xlsx')
def units_template(user=Depends(require_user)):
    wb=Workbook();sh=wb.active;sh.title='Units';sh.append(HEADERS)
    sh.append(['V-101','villa','available','B-01','BLD-1',0,'A',4,3,240,220,0,0,0,0,'detailed','',0,1800,450,600,0,2000,500,750,0,0,0,0,0])
    sh.append(['A-201','apartment','available','B-02','BLD-2',2,'B',2,2,110,0,12,4000,3600,300,'detailed','',0,2400,550,720,0,2650,600,900,0,0,0,0,0])
    sh.freeze_panes='A2';sh.auto_filter.ref=f'A1:AD3'
    for cell in sh[1]:
        from copy import copy
        style=copy(cell.font);style.bold=True;style.color='FFFFFF';cell.font=style
        cell.fill=__import__('openpyxl').styles.PatternFill('solid',fgColor='0B4332')
    for col in sh.columns:
        sh.column_dimensions[col[0].column_letter].width=min(32,max(14,len(str(col[0].value))+4))
    bio=io.BytesIO();wb.save(bio)
    return Response(bio.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':'attachment; filename="pricing_units_template.xlsx"'})

@app.post('/api/requests/{request_id}/import')
async def import_excel(request_id:int,file:UploadFile=File(...),user=Depends(require_user)):
    authorize(user,['maker'])
    if not file.filename or not file.filename.lower().endswith('.xlsx'):raise ValueError('يجب رفع ملف XLSX')
    payload=await file.read(6_000_001)
    if len(payload)>6_000_000:raise ValueError('الملف أكبر من 6MB')
    try:wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    except Exception:raise ValueError('ملف Excel غير صالح')
    ws=wb.active
    rows=ws.iter_rows(values_only=True)
    header=next(rows,None)
    if not header:raise ValueError('الملف فارغ')
    keys=[ALIASES.get(str(x).strip(),str(x).strip().lower()) if x is not None else '' for x in header]
    if not all(k in keys for k in REQUIRED_HEADERS):raise ValueError('أعمدة ملف الوحدات غير مكتملة؛ حمّل القالب من المنصة')
    prepared=[];errors=[];seen=set()
    for rownum,line in enumerate(rows,2):
        if not any(v is not None for v in line):continue
        if len(prepared)>=1000:raise ValueError('الحد الأقصى ألف وحدة لكل ملف')
        raw={k:v for k,v in zip(keys,line) if k in HEADERS and v is not None}
        code=str(raw.get('unit_code',''))
        if code in seen:errors.append(f'الصف {rownum}: رقم الوحدة مكرر');continue
        seen.add(code);prepared.append((rownum,raw))
    if not prepared:raise ValueError('لا توجد وحدات في الملف')
    # Atomic import: either all rows accepted or none (transaction rollback).
    with db(write=True) as c:
        req=must(c,'requests',request_id)
        if req['status'] not in ('draft','rework'):raise ValueError('الطلب غير قابل للاستيراد')
        if req['created_by']!=user['id'] and user['role']!='admin':raise HTTPException(403,'غير مسموح')
        p=must(c,'projects',req['project_id'])
        for rownum,raw in prepared:
            try:make_item(c,req,p,raw,'add' if req['kind']=='base' else 'update')
            except ValueError as ex:errors.append(f'الصف {rownum}: {ex}')
        if errors:raise ValueError('أخطاء الاستيراد: '+' | '.join(errors[:12]))
        audit(c,'request',request_id,request_id,user['id'],'excel_import',f'{file.filename}: {len(prepared)} rows')
    return {'imported':len(prepared)}

@app.post('/api/requests/{request_id}/preview')
def preview_request(request_id:int,user=Depends(require_user)):
    with db() as c:
        req=must(c,'requests',request_id);p=must(c,'projects',req['project_id'])
        warnings=check_request(c,req,p)
        rows=c.execute('SELECT * FROM request_items WHERE request_id=?',(request_id,)).fetchall()
        before=sum(D((jload(i['before_pricing']) or {}).get('segments',{}).get('beneficiary',{}).get('total',0)) for i in rows)
        after=sum(D((jload(i['after_pricing']) or {}).get('segments',{}).get('beneficiary',{}).get('total',0)) for i in rows)
        return {'count':len(rows),'before':money(before),'after':money(after),'delta':money(after-before),'warnings':warnings}

@app.post('/api/requests/{request_id}/action')
def request_action(request_id:int,p:dict=Body(...),user=Depends(require_user)):
    action=p.get('action');comment=str(p.get('comment','')).strip()
    with db(write=True) as c:
        req=must(c,'requests',request_id)
        if action=='submit':
            authorize(user,['maker'])
            if req['created_by']!=user['id'] and user['role']!='admin':raise HTTPException(403,'غير مسموح')
            if req['status'] not in ('draft','rework'):raise ValueError('الطلب ليس في حالة مسودة')
            check_request(c,req,must(c,'projects',req['project_id']))
            c.execute('UPDATE requests SET status="pending_checker",submitted_at=? WHERE id=?',(now(),request_id))
            c.execute('DELETE FROM approvals WHERE request_id=?',(request_id,))
            audit(c,'request',request_id,request_id,user['id'],'submitted')
            return {'status':'pending_checker'}
        if req['status']=='pending_checker':
            authorize(user,['checker'])
            if req['created_by']==user['id']:raise ValueError('لا يمكن أن يراجع منشئ الطلب طلبه')
            if action in ('rework','reject'):
                if not comment:raise ValueError('السبب مطلوب')
                status='rework' if action=='rework' else 'rejected'
                c.execute('UPDATE requests SET status=?,checked_by=? WHERE id=?',(status,user['id'],request_id))
                audit(c,'request',request_id,request_id,user['id'],action,comment)
                return {'status':status}
            if action!='approve':raise ValueError('إجراء غير متاح')
            check_request(c,req,must(c,'projects',req['project_id']))
            c.execute('UPDATE requests SET checked_by=? WHERE id=?',(user['id'],request_id))
            if req['kind']!='price':
                warnings=apply_request(c,req,user)
                return {'status':'approved','warnings':warnings}
            c.execute('UPDATE requests SET status="pending_approvals" WHERE id=?',(request_id,))
            for i,role in enumerate(APPROVAL_ROLES,1):
                c.execute('INSERT INTO approvals(request_id,step,role,status) VALUES(?,?,?,?)',
                          (request_id,i,role,'pending' if i==1 else 'waiting'))
            audit(c,'request',request_id,request_id,user['id'],'checker_approved',comment)
            return {'status':'pending_approvals'}
        if req['status']=='pending_approvals':
            step=c.execute('SELECT * FROM approvals WHERE request_id=? AND status="pending"',(request_id,)).fetchone()
            if not step:raise ValueError('لم يتم العثور على مرحلة الموافقة')
            authorize(user,[step['role']])
            if req['created_by']==user['id']:raise ValueError('لا يمكن للمُنشئ اعتماد طلبه')
            if action not in ('approve','rework','reject'):raise ValueError('إجراء غير متاح')
            if action in ('rework','reject') and not comment:raise ValueError('سبب الإرجاع أو الرفض مطلوب')
            result='approved' if action=='approve' else action
            c.execute('UPDATE approvals SET status=?,user_id=?,comment=?,acted_at=? WHERE id=?',
                      (result,user['id'],comment,now(),step['id']))
            audit(c,'request',request_id,request_id,user['id'],f'approval_{action}',f'{step["step"]} - {comment}')
            if action!='approve':
                status='rework' if action=='rework' else 'rejected'
                c.execute('UPDATE requests SET status=? WHERE id=?',(status,request_id))
                return {'status':status}
            next_step=c.execute('SELECT id FROM approvals WHERE request_id=? AND step=?',(request_id,step['step']+1)).fetchone()
            if next_step:
                c.execute('UPDATE approvals SET status="pending" WHERE id=?',(next_step['id'],))
                return {'status':'pending_approvals','next_step':step['step']+1}
            warnings=apply_request(c,req,user)
            return {'status':'approved','warnings':warnings}
        raise ValueError('هذا الطلب لا يقبل هذه العملية')

@app.get('/api/tasks')
def tasks(user=Depends(require_user)):
    with db() as c:
        if user['role']=='checker':
            rows=c.execute('''SELECT r.*,p.name project_name,u.name creator_name FROM requests r
                              JOIN projects p ON p.id=r.project_id JOIN users u ON u.id=r.created_by
                              WHERE status="pending_checker" AND r.created_by != ? ORDER BY r.id''',(user['id'],))
        elif user['role']=='admin':
            rows=c.execute('''SELECT r.*,p.name project_name,u.name creator_name FROM requests r
                              JOIN projects p ON p.id=r.project_id JOIN users u ON u.id=r.created_by
                              WHERE r.status IN ('pending_checker','pending_approvals') ORDER BY r.id''')
        else:
            rows=c.execute('''SELECT r.*,p.name project_name,u.name creator_name FROM requests r
                              JOIN approvals a ON a.request_id=r.id JOIN projects p ON p.id=r.project_id
                              JOIN users u ON u.id=r.created_by
                              WHERE r.status="pending_approvals" AND a.status="pending" AND a.role=? AND r.created_by!=?
                              ORDER BY r.id''',(user['role'],user['id']))
        return [dict(x) for x in rows]

@app.get('/api/compare')
def compare(a:int,b:int,user=Depends(require_user)):
    with db() as c:
        ra=must(c,'requests',a);rb=must(c,'requests',b)
        if ra['project_id']!=rb['project_id'] or ra['kind']!=rb['kind']:
            raise ValueError('تسمح المقارنة بطلبين من نفس المشروع ونفس النوع فقط')
        def indexed(req):
            return {i['unit_code']:dict(i) for i in c.execute('SELECT * FROM request_items WHERE request_id=?',(req['id'],))}
        aa=indexed(ra);bb=indexed(rb)
        differences=[]
        for code in sorted(set(aa)|set(bb)):
            xa=jload(aa[code]['after_json']) if code in aa else None
            xb=jload(bb[code]['after_json']) if code in bb else None
            keys=set((xa or {}).keys())|set((xb or {}).keys())
            diffs=[{'field':key,'a':(xa or {}).get(key),'b':(xb or {}).get(key)} for key in sorted(keys)
                   if (xa or {}).get(key)!=(xb or {}).get(key)]
            if diffs:differences.append({'unit_code':code,'fields':diffs})
        return {'a':ra['number'],'b':rb['number'],'matched':not differences,
                'differences':differences,'checked_units':len(set(aa)|set(bb))}

@app.get('/api/dashboard')
def dashboard(user=Depends(require_user)):
    with db() as c:
        stat={}
        for name,sql in [('destinations','SELECT COUNT(*) FROM destinations'),('projects','SELECT COUNT(*) FROM projects'),
                         ('units','SELECT COUNT(*) FROM units WHERE active=1'),('requests','SELECT COUNT(*) FROM requests'),
                         ('pending','SELECT COUNT(*) FROM requests WHERE status IN ("pending_checker","pending_approvals")'),
                         ('approved','SELECT COUNT(*) FROM requests WHERE status="approved"')]:
            stat[name]=c.execute(sql).fetchone()[0]
        stat['project_prices']=[]
        for x in c.execute('''SELECT p.id,p.name,COUNT(u.id) n,
                             COALESCE(SUM(CAST(json_extract(u.pricing_json,'$.segments.beneficiary.total') AS REAL)),0) amount
                             FROM projects p LEFT JOIN units u ON u.project_id=p.id AND u.active=1 GROUP BY p.id ORDER BY amount DESC LIMIT 8'''):
            stat['project_prices'].append(dict(x))
        stat['recent']=[dict(x) for x in c.execute('''SELECT r.id,r.number,r.kind,r.status,p.name project_name,r.created_at
                                 FROM requests r JOIN projects p ON p.id=r.project_id ORDER BY r.id DESC LIMIT 7''')]
        return stat

@app.get('/api/exports/units.xlsx')
def export_units(project_id:int,user=Depends(require_user)):
    with db() as c:
        must(c,'projects',project_id)
        units=[load_unit(x) for x in c.execute('SELECT * FROM units WHERE project_id=? AND active=1 ORDER BY unit_code',(project_id,))]
    wb=Workbook();sh=wb.active;sh.title='Units';sh.append(HEADERS+['beneficiary_total','nonbeneficiary_total'])
    for unit in units:
        data=unit['data'];prices=unit['pricing']['segments']
        sh.append([data.get(k,'') for k in HEADERS]+[prices['beneficiary']['total'],prices['nonbeneficiary']['total']])
    sh.freeze_panes='A2'
    bio=io.BytesIO();wb.save(bio)
    return Response(bio.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':f'attachment; filename="project_{project_id}_prices.xlsx"'})

@app.get('/api/reports')
def reports(user=Depends(require_user)):
    with db() as c:
        by_project=[dict(x) for x in c.execute('''SELECT p.name,p.code,COUNT(DISTINCT u.id) units,
          COALESCE(SUM(CAST(json_extract(u.pricing_json,'$.segments.beneficiary.total') AS REAL)),0) beneficiary_total,
          COALESCE(SUM(CAST(json_extract(u.pricing_json,'$.segments.nonbeneficiary.total') AS REAL)),0) nonbeneficiary_total
          FROM projects p LEFT JOIN units u ON u.project_id=p.id AND u.active=1 GROUP BY p.id ORDER BY p.id DESC''')]
        by_kind=[dict(x) for x in c.execute('SELECT kind,status,COUNT(*) count FROM requests GROUP BY kind,status')]
        return {'projects':by_project,'requests':by_kind}

@app.get('/api/health')
def health():return {'status':'ok','service':'pricing-portal'}

init_db()
